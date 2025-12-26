from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.http import HttpResponse, JsonResponse
from django.db.models import Count, Sum, Q
from datetime import datetime, timedelta, date
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import io

from .models import Client, Subscription, Visit, WriteOff, UserProfile, Reminder
from .forms import ClientForm, SubscriptionForm, VisitForm, UserCreateForm
from .decorators import admin_required, employee_required

def is_admin(user):
    try:
        return user.userprofile.role == 'admin'
    except:
        return False

def is_employee(user):
    return user.is_authenticated

@login_required
def dashboard(request):
    user = request.user
    today = date.today()
    
    # Для администратора
    if is_admin(user):
        context = {
            'today_visits': Visit.objects.filter(check_in__date=today).count(),
            'active_subscriptions': Subscription.objects.filter(status='active').count(),
            'expiring_soon': Subscription.objects.filter(
                status='active',
                end_date__range=[today, today + timedelta(days=7)]
            ).select_related('client'),
            'employee_stats': User.objects.filter(
                userprofile__role='employee'
            ).annotate(
                visit_count=Count('visits_created')
            ),
        }
    # Для сотрудника
    else:
        context = {
            'my_today_visits': Visit.objects.filter(
                created_by=user,
                check_in__date=today
            ).count(),
            'my_clients': Client.objects.filter(created_by=user).count(),
            'my_subscriptions': Subscription.objects.filter(
                created_by=user,
                status='active'
            ),
            'clients': Client.objects.filter(created_by=user)[:10],
        }
    
    return render(request, 'dashboard.html', context)

@login_required
def client_list(request):
    user = request.user
    if is_admin(user):
        clients = Client.objects.all()
    else:
        clients = Client.objects.filter(created_by=user)
    
    search = request.GET.get('search', '')
    if search:
        clients = clients.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(phone__icontains=search)
        )
    
    return render(request, 'clients/list.html', {'clients': clients, 'search': search})

@login_required
def client_create(request):
    if request.method == 'POST':
        form = ClientForm(request.POST)
        if form.is_valid():
            client = form.save(commit=False)
            client.created_by = request.user
            client.save()
            return redirect('client_list')
    else:
        form = ClientForm()
    
    return render(request, 'clients/form.html', {'form': form, 'title': 'Добавить клиента'})

@login_required
def client_detail(request, pk):
    client = get_object_or_404(Client, pk=pk)
    
    # Проверка прав доступа
    if not is_admin(request.user) and client.created_by != request.user:
        return redirect('dashboard')
    
    subscriptions = client.subscriptions.all()
    visits = client.visits.all()[:20]
    
    return render(request, 'clients/detail.html', {
        'client': client,
        'subscriptions': subscriptions,
        'visits': visits,
    })

@login_required
def client_edit(request, pk):
    client = get_object_or_404(Client, pk=pk)
    
    # Проверка прав доступа
    if not is_admin(request.user) and client.created_by != request.user:
        return redirect('dashboard')
    
    if request.method == 'POST':
        form = ClientForm(request.POST, instance=client)
        if form.is_valid():
            form.save()
            return redirect('client_detail', pk=pk)
    else:
        form = ClientForm(instance=client)
    
    return render(request, 'clients/form.html', {'form': form, 'title': 'Редактировать клиента', 'client': client})

@login_required
def client_delete(request, pk):
    client = get_object_or_404(Client, pk=pk)
    
    # Проверка прав доступа
    if not is_admin(request.user) and client.created_by != request.user:
        return redirect('dashboard')
    
    if request.method == 'POST':
        client.delete()
        return redirect('client_list')
    
    return render(request, 'clients/delete.html', {'client': client})

@login_required
def subscription_list(request):
    user = request.user
    if is_admin(user):
        subscriptions = Subscription.objects.all()
    else:
        subscriptions = Subscription.objects.filter(created_by=user)
    
    status = request.GET.get('status', '')
    if status:
        subscriptions = subscriptions.filter(status=status)
    
    return render(request, 'subscriptions/list.html', {
        'subscriptions': subscriptions,
        'status_filter': status,
    })

@login_required
def subscription_create(request):
    if request.method == 'POST':
        form = SubscriptionForm(request.POST)
        if form.is_valid():
            subscription = form.save(commit=False)
            subscription.created_by = request.user
            subscription.visits_left = subscription.visits_total
            subscription.save()
            return redirect('subscription_list')
    else:
        form = SubscriptionForm()
    
    # Ограничиваем выбор клиентов только своими (для сотрудников)
    if not is_admin(request.user):
        form.fields['client'].queryset = Client.objects.filter(created_by=request.user)
    
    return render(request, 'subscriptions/form.html', {'form': form, 'title': 'Создать абонемент'})

@login_required
def subscription_detail(request, pk):
    subscription = get_object_or_404(Subscription, pk=pk)
    
    # Проверка прав доступа
    if not is_admin(request.user) and subscription.created_by != request.user:
        return redirect('dashboard')
    
    writeoffs = subscription.writeoffs.all()
    visits = subscription.visits.all()
    
    return render(request, 'subscriptions/detail.html', {
        'subscription': subscription,
        'writeoffs': writeoffs,
        'visits': visits,
    })

@login_required
def visit_create(request):
    if request.method == 'POST':
        form = VisitForm(request.POST)
        if form.is_valid():
            visit = form.save(commit=False)
            visit.created_by = request.user
            
            # Автоматическое списание посещения
            if visit.subscription and visit.subscription.visits_left > 0:
                visit.subscription.visits_left -= 1
                if visit.subscription.visits_left == 0:
                    visit.subscription.status = 'expired'
                visit.subscription.save()
                
                # Создаем запись о списании
                WriteOff.objects.create(
                    subscription=visit.subscription,
                    visit=visit,
                    writeoff_type='auto',
                    amount=1,
                    created_by=request.user,
                    notes='Автоматическое списание при посещении'
                )
            
            visit.save()
            return redirect('dashboard')
    else:
        form = VisitForm()
    
    # Ограничиваем выбор клиентов и абонементов только своими (для сотрудников)
    if not is_admin(request.user):
        form.fields['client'].queryset = Client.objects.filter(created_by=request.user)
        form.fields['subscription'].queryset = Subscription.objects.filter(
            created_by=request.user,
            status='active',
            visits_left__gt=0
        )
    else:
        form.fields['subscription'].queryset = Subscription.objects.filter(
            status='active',
            visits_left__gt=0
        )
    
    return render(request, 'visits/create.html', {'form': form})

@login_required
def visit_checkout(request, pk):
    visit = get_object_or_404(Visit, pk=pk)
    visit.check_out = datetime.now()
    visit.save()
    return redirect('dashboard')

@login_required
@user_passes_test(is_admin)
def employee_list(request):
    employees = UserProfile.objects.filter(role='employee').select_related('user')
    return render(request, 'employees/list.html', {'employees': employees})

@login_required
@user_passes_test(is_admin)
def employee_create(request):
    if request.method == 'POST':
        form = UserCreateForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('employee_list')
    else:
        form = UserCreateForm()
    
    return render(request, 'employees/form.html', {'form': form, 'title': 'Добавить сотрудника'})

@login_required
def export_excel(request):
    user = request.user
    is_admin_user = is_admin(user)
    
    # Создаем Excel файл
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет"
    
    # Заголовки
    ws.append(['Отчет по посещениям'])
    ws.append(['Дата создания:', datetime.now().strftime('%Y-%m-%d %H:%M')])
    ws.append([''])
    
    # Данные
    if is_admin_user:
        visits = Visit.objects.all().select_related('client', 'created_by')
    else:
        visits = Visit.objects.filter(created_by=user).select_related('client')
    
    ws.append(['Клиент', 'Дата входа', 'Дата выхода', 'Продолжительность', 'Создано'])
    
    for visit in visits:
        duration = visit.duration() if visit.check_out else 'В зале'
        created_by = visit.created_by.username if visit.created_by else '-'
        ws.append([
            str(visit.client),
            visit.check_in.strftime('%Y-%m-%d %H:%M'),
            visit.check_out.strftime('%Y-%m-%d %H:%M') if visit.check_out else '-',
            duration,
            created_by
        ])
    
    # Стили
    for cell in ws[1]:
        cell.font = Font(bold=True, size=14)
    
    for cell in ws[4]:
        cell.font = Font(bold=True)
    
    # Создаем ответ
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="visits_report.xlsx"'
    
    wb.save(response)
    return response

@login_required
def export_pdf(request):
    user = request.user
    is_admin_user = is_admin(user)
    
    # Создаем PDF
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    
    styles = getSampleStyleSheet()
    elements.append(Paragraph("Отчет по посещениям", styles['Title']))
    elements.append(Paragraph(f"Дата создания: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Paragraph(" ", styles['Normal']))
    
    # Данные для таблицы
    if is_admin_user:
        visits = Visit.objects.all().select_related('client', 'created_by')[:50]
    else:
        visits = Visit.objects.filter(created_by=user).select_related('client')[:50]
    
    data = [['Клиент', 'Дата входа', 'Дата выхода', 'Продолжительность']]
    
    for visit in visits:
        duration = visit.duration() if visit.check_out else 'В зале'
        data.append([
            str(visit.client),
            visit.check_in.strftime('%Y-%m-%d %H:%M'),
            visit.check_out.strftime('%Y-%m-%d %H:%M') if visit.check_out else '-',
            duration
        ])
    
    # Создаем таблицу
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))
    
    elements.append(table)
    doc.build(elements)
    
    # Возвращаем PDF
    pdf = buffer.getvalue()
    buffer.close()
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="visits_report.pdf"'
    response.write(pdf)
    return response

@login_required
@user_passes_test(is_admin)
def admin_stats(request):
    today = date.today()
    week_ago = today - timedelta(days=7)
    
    stats = {
        'total_clients': Client.objects.count(),
        'total_subscriptions': Subscription.objects.count(),
        'active_subscriptions': Subscription.objects.filter(status='active').count(),
        'today_visits': Visit.objects.filter(check_in__date=today).count(),
        'week_visits': Visit.objects.filter(check_in__date__gte=week_ago).count(),
        'total_visits': Visit.objects.count(),
        'expiring_this_week': Subscription.objects.filter(
            status='active',
            end_date__range=[today, today + timedelta(days=7)]
        ).count(),
    }
    
    return render(request, 'admin/stats.html', {'stats': stats})